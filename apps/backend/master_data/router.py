from fastapi import APIRouter, Depends, Query, File, UploadFile, Request
from sqlalchemy.ext.asyncio import AsyncSession
from sqlalchemy import text, update, delete, insert
from core.database import get_db
from .service import process_nodeb_record
from pydantic import BaseModel
from typing import Optional, List
from datetime import date

router = APIRouter(prefix="/master-data")

class NodeBUpdate(BaseModel):
    sto: str
    site_id: str
    site_name: str
    hostname_metro: str
    ip_metro: Optional[str] = None
    port_metro: str
    hostname_olt: str
    ip_olt: Optional[str] = None
    type_olt: Optional[str] = None
    port_onu: Optional[str] = None
    hostname_ont: Optional[str] = None
    ip_ont: Optional[str] = None
    ont_type: Optional[str] = None
    serial_number: Optional[str] = None
    odc: Optional[str] = None
    odp: Optional[str] = None
    tikor_site: Optional[str] = None
    on_air: Optional[str] = None
    graph_id: Optional[str] = None
    evidence: Optional[str] = None

@router.get("/nodeb")
async def get_nodeb(
    request: Request,
    page: int = 1,
    limit: int = 25,
    search: str = "",
    fromdate: Optional[str] = Query(None),
    untildate: Optional[str] = Query(None),
    dateParam: Optional[str] = Query("on_air"),
    sort_by: Optional[str] = Query("idnodeb"),
    order: Optional[str] = Query("desc"),
    db: AsyncSession = Depends(get_db)
):
    offset = (page - 1) * limit
    
    # Parse query params manually for array support
    choice = request.query_params.getlist('choice[]') or request.query_params.getlist('choice')
    values = request.query_params.getlist('values[]') or request.query_params.getlist('values')
    select_fields = """
        n.idnodeb, n.idsto AS idsto_nodeb, n.site_id, n.site_name, 
        n.hostname_metro AS hostname_metro_nodeb, n.port_metro AS port_metro_nodeb,
        n.hostname_olt AS hostname_olt_nodeb, n.port_onu, n.ont_type,
        n.serial_number, n.odc, n.odp, n.tikor_site, n.on_air, n.hostname_ont, n.ip_ont,
        c.graph_id, ont.merk, ont.type,
        olt.port_metro AS port_metro_olt, olt.type_olt,
        sto.idsto, m.ip_metro, m.hostname_metro, olt.ip_olt
    """
    base_from = """
        FROM datanodeb n
        LEFT JOIN cacti c ON c.idnodeb = n.idnodeb
        LEFT JOIN dataont ont ON ont.serial_number = n.serial_number
        LEFT JOIN datasto sto ON sto.idsto = n.idsto
        LEFT JOIN datametro m ON m.hostname_metro = n.hostname_metro
        LEFT JOIN dataolt olt ON olt.hostname_olt = n.hostname_olt
        WHERE n.deleted_at IS NULL
    """
    where_clauses = []
    params = {"limit": limit, "offset": offset}

    if fromdate and untildate:
        db_date_col = "n.on_air" if dateParam == "on_air" else "n.created_at"
        where_clauses.append(f"{db_date_col} BETWEEN :fromdate AND :untildate")
        params["fromdate"] = f"{fromdate} 00:00:00"
        params["untildate"] = f"{untildate} 23:59:59"

    if search:
        where_clauses.append("(n.site_id LIKE :search OR n.site_name LIKE :search OR n.hostname_olt LIKE :search OR n.hostname_ont LIKE :search OR n.serial_number LIKE :search OR n.hostname_metro LIKE :search OR olt.ip_olt LIKE :search OR n.ip_ont LIKE :search OR n.odc LIKE :search OR n.odp LIKE :search OR n.ont_type LIKE :search)")
        params["search"] = f"%{search}%"

    if choice and values:
        choice_clauses = []
        for i, (c, v) in enumerate(zip(choice, values)):
            if not v or c == "cancel": continue
            col = c
            if c in ["idsto", "hostname_metro", "hostname_olt", "site_id", "site_name", "port_metro", "port_onu", "hostname_ont", "ip_ont", "ont_type", "serial_number", "odc", "odp", "tikor_site", "on_air"]:
                col = f"n.{c}"
            elif c == "ip_olt":
                col = "olt.ip_olt"
            elif c == "ip_metro":
                col = "m.ip_metro"
            p_key = f"val_{i}"
            choice_clauses.append(f"{col} LIKE :{p_key}")
            params[p_key] = f"%{v}%"
        if choice_clauses:
            where_clauses.append("(" + " OR ".join(choice_clauses) + ")")

    # Whitelist column mapping for sorting
    sort_mapping = {
        "idnodeb": "n.idnodeb",
        "idsto": "n.idsto",
        "site_id": "n.site_id",
        "site_name": "n.site_name",
        "hostname_metro": "n.hostname_metro",
        "port_metro": "n.port_metro",
        "hostname_olt": "n.hostname_olt",
        "ip_olt": "olt.ip_olt",
        "port_onu": "n.port_onu",
        "hostname_ont": "n.hostname_ont",
        "ip_ont": "n.ip_ont",
        "ont_type": "n.ont_type",
        "serial_number": "n.serial_number",
        "odc": "n.odc",
        "odp": "n.odp",
        "tikor_site": "n.tikor_site",
        "on_air": "n.on_air",
    }
    
    order_col = sort_mapping.get(sort_by, "n.idnodeb")
    order_dir = "ASC" if order.lower() == "asc" else "DESC"

    where_str = " AND " + " AND ".join(where_clauses) if where_clauses else ""
    count_query = f"SELECT COUNT(DISTINCT n.idnodeb) {base_from} {where_str}"
    data_query = f"SELECT {select_fields} {base_from} {where_str} GROUP BY n.idnodeb ORDER BY {order_col} {order_dir} LIMIT :limit OFFSET :offset"

    total = (await db.execute(text(count_query), params)).scalar() or 0
    result = await db.execute(text(data_query), params)

    return {
        "meta": {"total": total, "page": page, "limit": limit, "total_pages": (total + limit - 1) // limit if limit > 0 else 1},
        "data": [process_nodeb_record(dict(row)) for row in result.mappings().all()]
    }

@router.get("/nodeb/trash")
async def get_trash_nodeb(
    page: int = 1,
    limit: int = 25,
    search: str = "",
    fromdate: str = "",
    untildate: str = "",
    db: AsyncSession = Depends(get_db)
):
    offset = (page - 1) * limit
    base_from = """
        FROM datanodeb n
        LEFT JOIN cacti c ON c.idnodeb = n.idnodeb
        LEFT JOIN dataont ont ON ont.serial_number = n.serial_number
        LEFT JOIN datasto sto ON sto.idsto = n.idsto
        LEFT JOIN datametro m ON m.hostname_metro = n.hostname_metro
        LEFT JOIN dataolt olt ON olt.hostname_olt = n.hostname_olt
        WHERE n.deleted_at IS NOT NULL
    """
    count_query = f"SELECT COUNT(DISTINCT n.idnodeb) {base_from}"
    select_fields = """
        n.idnodeb, n.idsto AS idsto_nodeb, n.site_id, n.site_name, 
        n.hostname_metro AS hostname_metro_nodeb, n.port_metro AS port_metro_nodeb,
        n.hostname_olt AS hostname_olt_nodeb, olt.type_olt, n.port_onu, n.ont_type,
        n.serial_number, n.odc, n.odp, n.tikor_site, n.on_air, n.hostname_ont, n.ip_ont,
        c.graph_id, ont.merk, ont.type,
        olt.port_metro AS port_metro_olt, olt.ip_olt,
        sto.idsto,
        m.hostname_metro, n.deleted_at
    """
    data_query = f"SELECT {select_fields} {base_from}"
    if search:
        where_clause = """ AND (
            n.site_id LIKE :search OR n.site_name LIKE :search OR 
            n.hostname_olt LIKE :search OR n.hostname_ont LIKE :search OR n.serial_number LIKE :search
        )"""
        data_query += where_clause
        count_query += where_clause
        search_param = f"%{search}%"
    else:
        search_param = ""

    data_query += " GROUP BY n.idnodeb ORDER BY n.deleted_at DESC LIMIT :limit OFFSET :offset"

    if search:
        total = (await db.execute(text(count_query), {"search": search_param})).scalar() or 0
        result = await db.execute(text(data_query), {"search": search_param, "limit": limit, "offset": offset})
    else:
        total = (await db.execute(text(count_query))).scalar() or 0
        result = await db.execute(text(data_query), {"limit": limit, "offset": offset})

    return {
        "meta": {"total": total, "page": page, "limit": limit, "total_pages": (total + limit - 1) // limit if limit > 0 else 1},
        "data": [process_nodeb_record(dict(row)) for row in result.mappings().all()]
    }

from fastapi.responses import StreamingResponse
import io
import pandas as pd
import openpyxl.styles
import openpyxl.utils
from datetime import datetime, date

@router.get("/nodeb/export")
async def export_nodeb(
    request: Request,
    fromdate: Optional[str] = Query(None),
    untildate: Optional[str] = Query(None),
    dateParam: Optional[str] = Query("on_air"),
    db: AsyncSession = Depends(get_db)
):
    choice = request.query_params.getlist('choice[]') or request.query_params.getlist('choice')
    values = request.query_params.getlist('values[]') or request.query_params.getlist('values')
    select_fields = """
        n.idnodeb, n.idsto AS idsto_nodeb, n.site_id, n.site_name, 
        n.hostname_metro AS hostname_metro_nodeb, m.ip_metro, 
        olt.port_metro AS port_metro_olt, n.port_metro AS port_metro_nodeb,
        n.hostname_olt AS hostname_olt_nodeb, olt.ip_olt, 
        n.hostname_ont, n.ip_ont, n.port_onu, n.ont_type, 
        n.serial_number, n.odc, n.odp, n.tikor_site, n.on_air
    """
    from_clause = """
        FROM datanodeb n
        LEFT JOIN datasto sto ON sto.idsto = n.idsto
        LEFT JOIN datametro m ON m.hostname_metro = n.hostname_metro
        LEFT JOIN dataolt olt ON olt.hostname_olt = n.hostname_olt
        WHERE n.deleted_at IS NULL
    """
    
    where_clauses = []
    params = {}

    if fromdate and untildate:
        # Map frontend field names to DB column names if necessary
        db_date_col = "n.on_air" if dateParam == "on_air" else "n.created_at"
        where_clauses.append(f"{db_date_col} BETWEEN :fromdate AND :untildate")
        params["fromdate"] = f"{fromdate} 00:00:00"
        params["untildate"] = f"{untildate} 23:59:59"

    if choice and values:
        choice_clauses = []
        for i, (c, v) in enumerate(zip(choice, values)):
            if not v or c == "cancel": continue
            # Handle table prefixes for ambiguity
            col = c
            if c in ["idsto", "hostname_metro", "hostname_olt"]: col = f"n.{c}"
            elif c == "ip_olt": col = "olt.ip_olt"
            
            p_key = f"val_{i}"
            choice_clauses.append(f"{col} LIKE :{p_key}")
            params[p_key] = f"%{v}%"
        
        if choice_clauses:
            where_clauses.append("(" + " OR ".join(choice_clauses) + ")")

    where_str = ""
    if where_clauses:
        where_str = " AND " + " AND ".join(where_clauses)

    query = text(f"SELECT {select_fields} {from_clause} {where_str}")
    result = await db.execute(query, params)
    raw_data = result.mappings().all()

    processed_data = []
    for i, d in enumerate(raw_data):
        d_dict = dict(d)
        
        port_metro_val = d_dict['port_metro_nodeb']
        if d_dict['port_metro_olt'] is not None:
            port_metro_val = d_dict['port_metro_olt']

        processed_data.append({
            'NO': i + 1,
            'STO': d_dict['idsto_nodeb'],
            'SITE_ID': d_dict['site_id'],
            'SITE_NAME': d_dict['site_name'],
            'HOSTNAME_METRO': d_dict['hostname_metro_nodeb'],
            'IP_METRO': d_dict['ip_metro'],
            'PORT_METRO': port_metro_val,
            'HOSTNAME_OLT': d_dict['hostname_olt_nodeb'],
            'IP_OLT': d_dict['ip_olt'],
            'PORT_ONU': d_dict['port_onu'],
            'HOSTNAME_ONT': d_dict['hostname_ont'],
            'IP_ONT': d_dict['ip_ont'],
            'ONT_TYPE': d_dict['ont_type'],
            'SERIAL_NUMBER': d_dict['serial_number'],
            'ODC': d_dict['odc'],
            'ODP': d_dict['odp'],
            'COORDINATE': d_dict['tikor_site'],
            'ON AIR': d_dict['on_air'],
        })

    df = pd.DataFrame(processed_data)
    stream = io.BytesIO()

    writer = pd.ExcelWriter(stream, engine='openpyxl')
    df.to_excel(writer, index=False, sheet_name='NodeB Data')
    
    workbook = writer.book
    worksheet = writer.sheets['NodeB Data']

    header_fill = openpyxl.styles.PatternFill(start_color="499ff2", end_color="499ff2", fill_type="solid")
    header_font = openpyxl.styles.Font(bold=True, color="FFFFFF")
    
    thin_border = openpyxl.styles.Border(left=openpyxl.styles.Side(style='thin', color='FF000000'),
                                        right=openpyxl.styles.Side(style='thin', color='FF000000'),
                                        top=openpyxl.styles.Side(style='thin', color='FF000000'),
                                        bottom=openpyxl.styles.Side(style='thin', color='FF000000'))
    
    for col_idx in range(1, len(df.columns) + 1):
        cell = worksheet.cell(row=1, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font
        worksheet.column_dimensions[openpyxl.utils.get_column_letter(col_idx)].auto_size = True
        
    for row_idx in range(1, len(df) + 2):
        for col_idx in range(1, len(df.columns) + 1):
            cell = worksheet.cell(row=row_idx, column=col_idx)
            cell.border = thin_border

    writer.close()
    stream.seek(0)
    
    filename = f'NODEB-{datetime.now().strftime("%y%m%d-%H%M%S")}.xlsx'
    
    return StreamingResponse(
        stream,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename=\"{filename}\""}
    )

@router.post("/nodeb/import")
async def import_nodeb(file_excel: UploadFile = File(...), db: AsyncSession = Depends(get_db)):
    return {"status": "success"}

@router.get("/nodeb/{id}")
async def get_nodeb_detail(id: int, db: AsyncSession = Depends(get_db)):
    query = text("""
        SELECT 
            n.idnodeb, n.idsto AS idsto_nodeb, n.site_id, n.site_name, 
            n.hostname_metro AS hostname_metro_nodeb, n.port_metro AS port_metro_nodeb,
            n.hostname_olt AS hostname_olt_nodeb, n.port_onu, n.ont_type,
            n.serial_number, n.odc, n.odp, n.tikor_site, n.on_air, n.hostname_ont, n.ip_ont,
            c.graph_id, ont.merk, ont.type,
            olt.port_metro AS port_metro_olt, olt.type_olt,
            sto.idsto, m.ip_metro, m.hostname_metro, olt.ip_olt
        FROM datanodeb n
        LEFT JOIN cacti c ON c.idnodeb = n.idnodeb
        LEFT JOIN dataont ont ON ont.serial_number = n.serial_number
        LEFT JOIN datasto sto ON sto.idsto = n.idsto
        LEFT JOIN datametro m ON m.hostname_metro = n.hostname_metro
        LEFT JOIN dataolt olt ON olt.hostname_olt = n.hostname_olt
        WHERE n.idnodeb = :id
    """)
    result = await db.execute(query, {"id": id})
    row = result.mappings().first()
    if not row:
        return {"error": "Not found"}
    return process_nodeb_record(dict(row))

@router.post("/nodeb")
async def create_nodeb(data: NodeBUpdate, db: AsyncSession = Depends(get_db)):
    insert_data = data.dict(exclude={'graph_id', 'evidence'})
    stmt = insert(text("datanodeb")).values(**insert_data)
    result = await db.execute(stmt)
    new_id = result.inserted_primary_key[0]
    
    if data.graph_id:
        await db.execute(text("INSERT INTO cacti (graph_id, idnodeb) VALUES (:g, :id)"), {"g": data.graph_id, "id": new_id})
        
    await db.commit()
    return {"status": "success", "id": new_id}

@router.patch("/nodeb/{id}")
async def update_nodeb(id: int, data: NodeBUpdate, db: AsyncSession = Depends(get_db)):
    update_data = data.dict(exclude={'graph_id', 'evidence'})
    stmt = update(text("datanodeb")).where(text("idnodeb = :id")).values(**update_data)
    await db.execute(stmt, {"id": id})
    
    if data.graph_id is not None:
        await db.execute(delete(text("cacti")).where(text("idnodeb = :id")), {"id": id})
        if data.graph_id:
            await db.execute(text("INSERT INTO cacti (graph_id, idnodeb) VALUES (:g, :id)"), {"g": data.graph_id, "id": id})
            
    await db.commit()
    return {"status": "success"}

@router.post("/nodeb/restore/{id}")
async def restore_nodeb(id: int, db: AsyncSession = Depends(get_db)):
    await db.execute(text("UPDATE datanodeb SET deleted_at = NULL WHERE idnodeb = :id"), {"id": id})
    await db.commit()
    return {"status": "success"}

@router.post("/nodeb/restore-all")
async def restore_all_nodeb(db: AsyncSession = Depends(get_db)):
    await db.execute(text("UPDATE datanodeb SET deleted_at = NULL WHERE deleted_at IS NOT NULL"))
    await db.commit()
    return {"status": "success"}

@router.delete("/nodeb/{id}")
async def delete_nodeb(id: int, db: AsyncSession = Depends(get_db)):
    await db.execute(text("UPDATE datanodeb SET deleted_at = datetime('now') WHERE idnodeb = :id"), {"id": id})
    await db.commit()
    return {"status": "success"}

@router.delete("/nodeb/purge/{id}")
async def purge_nodeb(id: int, db: AsyncSession = Depends(get_db)):
    await db.execute(text("DELETE FROM datanodeb WHERE idnodeb = :id"), {"id": id})
    await db.commit()
    return {"status": "success"}

@router.delete("/nodeb/purge-all")
async def purge_all_nodeb(db: AsyncSession = Depends(get_db)):
    await db.execute(text("DELETE FROM datanodeb WHERE deleted_at IS NOT NULL"))
    await db.commit()
    return {"status": "success"}

@router.get("/sto")
async def get_sto(db: AsyncSession = Depends(get_db)):
    result = await db.execute(text("SELECT idsto FROM datasto WHERE idsto != '-'"))
    return [r[0] for r in result.fetchall()]

@router.get("/olt")
async def get_olt(
    page: int = 1,
    limit: int = 10,
    search: str = "",
    db: AsyncSession = Depends(get_db)
):
    offset = (page - 1) * limit
    base_from = "FROM dataolt WHERE hostname_olt != '-'"
    if search:
        where = " AND (hostname_olt LIKE :s OR ip_olt LIKE :s OR platform LIKE :s)"
        search_p = f"%{search}%"
    else:
        where = ""
        search_p = ""

    count_q = f"SELECT COUNT(*) {base_from} {where}"
    data_q = f"SELECT hostname_metro, port_metro, hostname_olt, ip_olt, platform {base_from} {where} LIMIT :limit OFFSET :offset"

    if search:
        total = (await db.execute(text(count_q), {"s": search_p})).scalar() or 0
        res = await db.execute(text(data_q), {"s": search_p, "limit": limit, "offset": offset})
    else:
        total = (await db.execute(text(count_q))).scalar() or 0
        res = await db.execute(text(data_q), {"limit": limit, "offset": offset})

    return {
        "meta": {"total": total, "page": page, "limit": limit, "total_pages": (total + limit - 1) // limit if limit > 0 else 1},
        "data": [dict(r) for r in res.mappings().all()]
    }
